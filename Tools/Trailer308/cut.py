"""트레일러 편집 — ffmpeg 한 번의 -filter_complex_script 로 컷을 잇고, audio.py의 WAV를 입혀 Art/Pitch308/Trailer/trailer308_<N>_v4.mp4.

컷 체인(컷마다): [k:v]trim=duration=dur·speed+0.2, setpts=(PTS-STARTPTS)/speed, fps=30, (scale), crop=1920:804:0:138,
  pad=1920:1080:0:138:color=black, (grade), tpad=stop_mode=clone:stop_duration=dur, trim=end_frame=N, setpts=PTS-STARTPTS, (fade dip), format=yuv420p
  → concat → -frames:v N총 -r 30 -c:v libx264 -preset medium -crf 17 -pix_fmt yuv420p -movflags +faststart.
  (tpad 는 컷 길이만큼 마지막 프레임을 복제한다 — 테이크가 창보다 짧거나 held_last 로 늘어나도 프레임이 모자라지 않는다. 비용 없음.)
첫 컷 inkin = title.py의 PIL 마스크 27장(0.9 s) + maskedmerge(gbrp). 입력은 전부 `-ss <src> -t <need> -i <path>` — 필터 문자열에 경로 · 폰트 없음.

usage: python cut.py 30|90 [--preview] [--dry-run] [--fake] [--music X|none] [--music-offset S] [--remux] [--no-audio]
                           [--grade none|blackpoint] [--inkin mask|fade] [--deck] [--make-test-takes] [--force-title]
  --preview          960×540 · ultrafast · crf 23 → _work/preview_<N>.mp4 (보고서는 _work/report_<N>_preview.json)
  --fake             _work/test_takes/<take>.mp4(testsrc2 카운터, 20 s)를 입력으로. 없으면 만든다
  --remux            영상 재렌더 없이 _work/video_<N>.mp4 + 새 WAV → 최종. 컷은 Films 를 다시 보지 않고 **report_<N>.json 의 cuts** 를 쓴다
                     (렌더 뒤 Films 가 바뀌어도 소리 marks · 엔벌로프가 영상과 어긋나지 않는다). report 가 없거나 영상과 안 맞으면 중단.
  --deck             최종 → trailer308_deck_v4.mp4(≤ 25 MB 규격) + 제출본 pptx **사본** _work/deck_test.pptx에 add_movie 시험
보고서: _work/report_<N>.json (쓴 컷 · ALT · FREEZE · 명령 · 렌더 소요 실측). --remux 는 cuts · resolve · inputs · video 를 덮어쓰지 않고
  remux 항목(시각 · argv · 소리 · 최종)만 덧붙인다.
"""
import argparse
import hashlib
import shutil
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (BAR, DECK_OUT, DECK_PPTX, FFMPEG, FILMS, FPS, H, OUT, OUTPUT, VIEW_H, W, WORK, count_frames,  # noqa: E402
                    ensure_dirs, fmt_cmd, probe, read_json, run, scan_films, write_json)
from timeline import AbortTimeline, Cut, all_takes, describe, resolve, starts, timeline, validate  # noqa: E402

TEST_TAKES = WORK / 'test_takes'
TEST_TAKE_SECONDS = 20
INKIN_SECONDS = 0.9
DIP = 0.12


def report_path(N, preview=False):
    return WORK / f'report_{N}{"_preview" if preview else ""}.json'


def load_report_for_remux(N, preview, video):
    """--remux: report_<N>[_preview].json 의 cuts 를 정본으로 쓴다. 없거나 영상과 안 맞으면 (None, 이유)."""
    rp = report_path(N, preview)
    rep = read_json(rp)
    if not rep or not rep.get('cuts'):
        return None, f'--remux: 보고서가 없다 {rp} — 먼저 렌더(cut.py {N}{" --preview" if preview else ""})'
    if not video.exists():
        return None, f'--remux: 영상이 없다 {video}'
    cuts = [Cut.from_dict(x) for x in rep['cuts']]
    total = sum(c.frames for c in cuts)
    got = count_frames(video)
    if got != total:
        return None, f'--remux: 영상 {video.name} {got} f ≠ 보고서 컷 합 {total} f — 보고서와 영상이 다른 렌더다. 다시 렌더하라'
    return (rep, cuts), None


# ---------------------------------------------------------------- 가짜 테이크
def make_test_takes(force=False):
    """testsrc2 1920×1080 30 fps 20 s. 카운터(왼쪽 위 0–35행)가 138행 크롭에 잘리지 않게 패턴을 138 px 내리고 띠는 회색(0x404040)으로
    채운다 — 출력에 회색이 보이면 크롭이 어긋난 것. 테이크마다 hue를 돌려 어느 테이크인지 눈으로 구분한다."""
    TEST_TAKES.mkdir(parents=True, exist_ok=True)
    made = []
    for i, take in enumerate(all_takes()):
        out = TEST_TAKES / f'{take}.mp4'
        if out.exists() and not force:
            continue
        vf = f'crop={W}:{VIEW_H}:0:0,pad={W}:{H}:0:{BAR}:color=0x404040,hue=h={(i * 37) % 360}'
        run([FFMPEG, '-y', '-loglevel', 'error', '-f', 'lavfi', '-i', f'testsrc2=size={W}x{H}:rate={FPS}', '-t', str(TEST_TAKE_SECONDS),
             '-vf', vf, '-c:v', 'libx264', '-preset', 'veryfast', '-crf', '12', '-pix_fmt', 'yuv420p', '-movflags', '+faststart', out])
        made.append(out)
    return made


# ---------------------------------------------------------------- 필터 스크립트
def build_filter(cuts, films, title_idx, mask_idx, opts):
    """(inputs argv, 스크립트 문자열, 입력표). 입력 k = 샷 컷 순서, 그다음 제목, 그다음 마스크."""
    inputs, table, lines = [], [], []
    k = 0
    for i, c in enumerate(cuts):
        N = c.frames
        if c.kind == 'shot':
            info = films[c.take]
            path = info['path']
            need = c.dur * c.speed + 0.2
            inputs += ['-ss', f'{c.src:.4f}', '-t', f'{need + 0.1:.4f}', '-i', path]
            table.append({'input': k, 'name': c.name, 'take': c.take, 'path': path, 'src': c.src, 't': round(need + 0.1, 4)})
            parts = [f'[{k}:v]trim=duration={need:.4f}', f'setpts=(PTS-STARTPTS)/{c.speed:g}', f'fps={FPS}']
            if (info.get('width'), info.get('height')) != (W, H):
                parts.append(f'scale={W}:{H}:flags=bicubic')
            parts += [f'crop={W}:{VIEW_H}:0:{BAR}', f'pad={W}:{H}:0:{BAR}:color=black']
            if opts.grade == 'blackpoint':
                parts.append('colorlevels=rimin=0.02:gimin=0.02:bimin=0.02')
            parts += [f'tpad=stop_mode=clone:stop_duration={c.dur:.3f}', f'trim=end_frame={N}', 'setpts=PTS-STARTPTS']
            tail = []
            if c.fx_out == 'dip':
                tail.append(f'fade=t=out:st={c.dur - DIP:.4f}:d={DIP}')
            if c.fx_in == 'inkin' and opts.inkin == 'mask' and mask_idx is not None:
                lines.append(','.join(parts) + f',format=gbrp[c{i}]')
                lines.append(f'color=c=black:s={W}x{H}:r={FPS}:d=1,format=gbrp[blk{i}]')
                lines.append(f'[{mask_idx}:v]format=gbrp[msk{i}]')
                lines.append(f'[blk{i}][c{i}][msk{i}]maskedmerge,trim=end_frame={N},setpts=PTS-STARTPTS' + ''.join(',' + t for t in tail) + f',format=yuv420p[v{i}]')
            else:
                if c.fx_in == 'inkin':
                    tail.insert(0, f'fade=t=in:st=0:d={INKIN_SECONDS}')
                lines.append(','.join(parts + tail) + f',format=yuv420p[v{i}]')
            k += 1
        elif c.kind == 'black':
            lines.append(f'color=c=black:s={W}x{H}:r={FPS}:d={c.dur + 0.2:.3f},trim=end_frame={N},setpts=PTS-STARTPTS,format=yuv420p[v{i}]')
        elif c.kind == 'title':
            lines.append(f'[{title_idx}:v]fps={FPS},tpad=stop_mode=clone:stop_duration=1,trim=end_frame={N},setpts=PTS-STARTPTS,format=yuv420p[v{i}]')
        else:
            raise ValueError(c.kind)
    labels = ''.join(f'[v{i}]' for i in range(len(cuts)))
    lines.append(f'{labels}concat=n={len(cuts)}:v=1:a=0[vcat]')
    if opts.preview:
        lines.append('[vcat]scale=960:540:flags=bicubic,format=yuv420p[vout]')
    else:
        lines.append('[vcat]format=yuv420p[vout]')
    return inputs, ';\n'.join(lines) + '\n', table


# ---------------------------------------------------------------- 덱
def sha1(path):
    h = hashlib.sha1()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def encode_deck(final, deck, log):
    cmd = [FFMPEG, '-y', '-hide_banner', '-loglevel', 'warning', '-i', final, '-c:v', 'libx264', '-profile:v', 'high', '-level', '4.0',
           '-preset', 'slow', '-crf', '23', '-maxrate', '6M', '-bufsize', '12M', '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-b:a', '128k',
           '-ar', '48000', '-ac', '2', '-movflags', '+faststart', deck]
    _, dt = run(cmd, log=log)
    return cmd, dt


def pptx_test(deck, log):
    """제출본의 사본 _work/deck_test.pptx 에 add_movie. 원본은 해시로 불변 확인."""
    from pptx import Presentation
    from pptx.util import Emu
    if not DECK_PPTX.exists():
        return {'skipped': f'pptx 없음: {DECK_PPTX}'}
    before = sha1(DECK_PPTX)
    test = WORK / 'deck_test.pptx'
    shutil.copyfile(DECK_PPTX, test)
    poster = WORK / 'deck_poster.png'
    run([FFMPEG, '-y', '-loglevel', 'error', '-ss', '20', '-i', deck, '-frames:v', '1', poster], log=log)
    prs = Presentation(str(test))
    blank = next((l for l in prs.slide_layouts if l.name.lower() == 'blank'), prs.slide_layouts[-1])
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_movie(str(deck), Emu(0), Emu(0), prs.slide_width, prs.slide_height, poster_frame_image=str(poster), mime_type='video/mp4')
    prs.save(str(test))
    after = sha1(DECK_PPTX)
    return {'pptx_copy': str(test), 'copy_bytes': test.stat().st_size, 'slides': len(prs.slides), 'original_unchanged': before == after,
            'original_sha1': before}


# ---------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('length', type=int, choices=(30, 90))
    ap.add_argument('--preview', action='store_true')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--fake', action='store_true')
    ap.add_argument('--music', default=None)
    ap.add_argument('--music-offset', type=float, default=0.0)
    ap.add_argument('--remux', action='store_true')
    ap.add_argument('--no-audio', action='store_true')
    ap.add_argument('--grade', choices=('none', 'blackpoint'), default='none')
    ap.add_argument('--inkin', choices=('mask', 'fade'), default='mask')
    ap.add_argument('--deck', action='store_true')
    ap.add_argument('--make-test-takes', action='store_true')
    ap.add_argument('--force-title', action='store_true')
    ap.add_argument('--films', default=None, help='테이크 폴더 덮어쓰기(시험용 — 기본 Films/, --fake 면 _work/test_takes/)')
    a = ap.parse_args(argv)
    ensure_dirs()
    N = a.length
    log = WORK / f'log_cut_{N}.txt'
    rpath = report_path(N, a.preview)
    video = WORK / (f'preview_{N}_video.mp4' if a.preview else f'video_{N}.mp4')
    final = (WORK / f'preview_{N}.mp4') if a.preview else (OUT / OUTPUT[N])
    script = WORK / f'filter_{N}{"_preview" if a.preview else ""}.txt'

    if a.make_test_takes:
        t = time.perf_counter()
        made = make_test_takes()
        print(f'test takes: {len(made)} made, {len(list(TEST_TAKES.glob("*.mp4")))} present  ({time.perf_counter() - t:.2f}s)')
        if not a.fake:
            return

    if a.remux:
        # 컷의 정본 = 렌더 때의 보고서. Films 를 다시 resolve 하지 않는다(렌더 뒤 Films 가 바뀌면 소리가 영상과 어긋난다 — 검토 지적).
        loaded, why = load_report_for_remux(N, a.preview, video)
        if loaded is None:
            print('ERROR', why)
            sys.exit(2)
        report, cuts = loaded
        if bool(report.get('fake')) != a.fake:
            print(f"WARN --remux: 보고서는 fake={report.get('fake')} 로 렌더됐다(지금 인자 --fake={a.fake}) — 보고서의 컷을 쓴다")
        total = sum(c.frames for c in cuts)
        print(describe(cuts, N))
        print(f'remux ← {rpath.name} cuts (rendered {report.get("video", {}).get("frames")} f, {report.get("mode")})')
        report.setdefault('remux_history', []).append({'at': time.strftime('%Y-%m-%d %H:%M:%S'), 'argv': sys.argv[1:]})
        report['commands'] = [c for c in report.get('commands', []) if c.get('step') == 'render']   # 렌더 명령은 남기고 소리 · 덱은 새로
        report.setdefault('timing', {})
        for k in ('audio', 'mux', 'deck', 'deck_total'):
            report['timing'].pop(k, None)
    else:
        report = {'length': N, 'mode': 'preview' if a.preview else 'final', 'fake': a.fake, 'argv': sys.argv[1:], 'commands': [], 'timing': {}}
        films_dir = Path(a.films) if a.films else (TEST_TAKES if a.fake else FILMS)
        if a.fake and not a.films and not any(TEST_TAKES.glob('*.mp4')):
            make_test_takes()
        films = scan_films(films_dir)
        try:
            cuts, res = resolve(timeline(N), {k: v.get('duration') for k, v in films.items()})
        except AbortTimeline as e:
            print('ABORT', e)
            sys.exit(2)
        errors, warnings = validate(cuts, N, {k: v.get('duration') for k, v in films.items()})
        print(describe(cuts, N))
        for w in warnings:
            print('WARN', w)
        if errors:
            for e in errors:
                print('ERROR', e)
            sys.exit(2)
        for x in res['alt']:
            print(f"ALT  {x['name']}: {x['take']} → {x['alt']}")
        for x in res['dropped']:
            print(f"DROP {x['name']} ({x['take']}, {x['dur']}s)")
        for x in res['freeze']:
            print(f"FREEZE {x['name']} ({x['take']}) +{x['frames']}f clone")
        if res['held_last']:
            print(f"HELD {res['held_last']['name']} +{res['held_last']['added_frames']}f → {res['held_last']['dur']}s")
        total = sum(c.frames for c in cuts)
        t0s = starts(cuts)
        report.update({'cuts': [dict(c.to_dict(), t0=t0) for c, t0 in zip(cuts, t0s)], 'resolve': res, 'frames': total,
                       'play_seconds': res['play_seconds'], 'films_dir': str(films_dir)})

    if not a.remux:
        import title as title_mod
        t = time.perf_counter()
        title_mp4 = title_mod.render_title(N, next(c.dur for c in cuts if c.kind == 'title'), force=a.force_title)
        mask_dir = title_mod.write_inkin() if a.inkin == 'mask' else None
        report['timing']['title'] = round(time.perf_counter() - t, 2)
        n_shots = sum(1 for c in cuts if c.kind == 'shot')
        title_idx = n_shots
        mask_idx = n_shots + 1 if mask_dir else None
        inputs, text, table = build_filter(cuts, films, title_idx, mask_idx, a)
        inputs += ['-i', str(title_mp4)]
        if mask_dir:
            inputs += ['-framerate', str(FPS), '-i', str(mask_dir / 'ink_%02d.png')]
        script.write_text(text, encoding='utf-8', newline='\n')
        preset, crf = ('ultrafast', '23') if a.preview else ('medium', '17')
        cmd = [FFMPEG, '-y', '-hide_banner', '-loglevel', 'warning', '-stats'] + inputs + \
              ['-filter_complex_script', str(script), '-map', '[vout]', '-frames:v', str(total), '-r', str(FPS), '-c:v', 'libx264',
               '-preset', preset, '-crf', crf, '-pix_fmt', 'yuv420p', '-movflags', '+faststart', '-an', str(video)]
        report['inputs'] = table
        report['filter_script'] = str(script)
        report['commands'].append({'step': 'render', 'cmd': fmt_cmd(cmd)})
        if a.dry_run:
            print('--- filter script', script)
            print(text)
            print('--- command')
            print(fmt_cmd(cmd))
            write_json(rpath, report)
            return
        print(f'render → {video}')
        _, dt = run(cmd, log=log)
        got = count_frames(video)
        report['timing']['render'] = round(dt, 2)
        report['video'] = {'path': str(video), 'frames': got, 'bytes': video.stat().st_size}
        print(f'rendered {got} frames in {dt:.1f}s ({total / max(dt, 1e-6):.1f} f/s)  {video.stat().st_size / 1e6:.1f} MB')
        if got != total:
            print(f'ERROR frames {got} != {total}')
            write_json(rpath, report)
            sys.exit(1)

    if a.no_audio:
        cmd = [FFMPEG, '-y', '-loglevel', 'error', '-i', video, '-c:v', 'copy', '-an', '-movflags', '+faststart', final]
        _, dt = run(cmd, log=log)
        report['commands'].append({'step': 'copy', 'cmd': fmt_cmd(cmd)})
    else:
        import audio as audio_mod
        t = time.perf_counter()
        music = audio_mod.music_path(a.music, N)
        wav = WORK / (f'audio_{N}{"_preview" if a.preview else ""}.wav')
        ainfo = audio_mod.build(N, cuts, music, a.music_offset, wav)
        report['timing']['audio'] = round(time.perf_counter() - t, 2)
        report['audio'] = {k: v for k, v in ainfo.items() if k != 'sfx'}
        print(f"audio {wav.name} {ainfo['seconds']:.3f}s peak {ainfo['peak_after']:.3f} sfx {ainfo['sfx_count']} music {music}")
        cmd = [FFMPEG, '-y', '-loglevel', 'error', '-i', video, '-i', wav, '-map', '0:v', '-map', '1:a', '-c:v', 'copy', '-c:a', 'aac',
               '-b:a', '192k', '-ar', '48000', '-ac', '2', '-movflags', '+faststart', final]
        _, dt = run(cmd, log=log)
        report['timing']['mux'] = round(dt, 2)
        report['commands'].append({'step': 'mux', 'cmd': fmt_cmd(cmd)})
    fin = probe(final)
    report['final'] = {'path': str(final), 'bytes': fin['bytes'], 'duration': fin['duration'], 'frames': count_frames(final),
                       'size': f"{fin['width']}x{fin['height']}", 'fps': fin['fps'], 'audio': f"{fin['acodec']} {fin['sample_rate']} {fin['channels']}"}
    print(f"final {final}  {fin['bytes'] / 1e6:.1f} MB  {report['final']['frames']} f  {report['final']['size']}  {report['final']['audio']}")

    if a.deck:
        deck = OUT / DECK_OUT[N]
        t = time.perf_counter()
        cmd, dt = encode_deck(final, deck, log)
        report['commands'].append({'step': 'deck', 'cmd': fmt_cmd(cmd)})
        report['timing']['deck'] = round(dt, 2)
        d = probe(deck)
        report['deck'] = {'path': str(deck), 'bytes': d['bytes'], 'frames': count_frames(deck)}
        print(f"deck {deck}  {d['bytes'] / 1e6:.1f} MB  {report['deck']['frames']} f  ({dt:.1f}s)")
        try:
            report['deck']['pptx'] = pptx_test(deck, log)
            print('pptx', report['deck']['pptx'])
        except Exception as e:                           # noqa: BLE001 — 덱 삽입 시험 실패는 보고서에만
            report['deck']['pptx'] = {'error': str(e)}
            print('pptx ERROR', e)
        report['timing']['deck_total'] = round(time.perf_counter() - t, 2)
    write_json(rpath, report)
    print('report', rpath, 'timing', report['timing'])


if __name__ == '__main__':
    main()
