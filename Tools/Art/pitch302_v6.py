"""#302 pitch deck v6: minimal, dark page, white type. One layout spec rendered as an HTML board and as the .pptx.

Page 먹 111112, type F2F0EC, secondary 8E8A84, hairlines 2E2D2B, one accent 황토 D9C39A used once or twice per page.
One family (KoPub 돋움: Bold / Medium / Light). No cards, no boxes: hairlines and space divide the page.
Canvas 1280 x 720 px = 13.333 x 7.5 in (96 px/in). Charts are native (the margin curve is an XY chart so the 10,000-copy
drop line is a real dashed series).

usage: python Tools/Art/pitch302_v6.py [--trailer path.mp4] [--out name.pptx] [--no-notes] [--no-embed] [--html-only]
Output: Art/Presentation302/v6/board.html, Art/Presentation302/오행부_피칭_v6.pptx
"""
import html
import math
import subprocess
import sys
from pathlib import Path

import imageio_ffmpeg
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'Art/Presentation302'
BOARD = OUT / 'v6'
SRC = OUT / 'src'
SHOTS = ROOT / 'Art/Presentation297/Screenshots'
STILLS = ROOT / 'Art/Presentation297/Stills'
TRAILER = Path(sys.argv[sys.argv.index('--trailer') + 1]) if '--trailer' in sys.argv else ROOT / 'Art/Presentation297/Trailer/오행부_트레일러_1분_드라이브용.mp4'
DECK = OUT / (sys.argv[sys.argv.index('--out') + 1] if '--out' in sys.argv else '오행부_피칭_v6.pptx')   # --out <file name>

BG, TEXT, MUTED, RULE, ACCENT, DIM = '111112', 'F2F0EC', '8E8A84', '2E2D2B', 'D9C39A', 'BDB9B2'
B, M, LT = 'KoPubDotum Bold', 'KoPubDotum Medium', 'KoPubDotum Light'
FILES = {B: 'KoPubDotumBold', M: 'KoPubDotumMedium', LT: 'KoPubDotumLight'}
W, H = 1280, 720
X0, X1 = 72, 1208
TOP = 150


# ---------------------------------------------------------------- layout model

def T(x, y, w, h, runs, size=15, font=M, colour=TEXT, align='left', lh=1.5):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, font, colour)."""
    return dict(k='text', x=x, y=y, w=w, h=h, runs=runs, size=size, font=font, colour=colour, align=align, lh=lh)


def R(x, y, w, h, fill):
    return dict(k='rect', x=x, y=y, w=w, h=h, fill=fill)


def HR(x0, x1, y, colour=RULE):
    return R(x0, y, x1 - x0, 1, colour)


def VR(x, y0, y1, colour=RULE):
    return R(x, y0, 1, y1 - y0, colour)


def I(path, x, y, w, h):
    return dict(k='img', path=Path(path), x=x, y=y, w=w, h=h)


def CHART(kind, x, y, w, h, **data):
    return dict(k='chart', kind=kind, x=x, y=y, w=w, h=h, **data)


def VIDEO(path, poster, x, y, w, h):
    return dict(k='video', path=Path(path), poster=Path(poster), x=x, y=y, w=w, h=h)


def page(n, title):
    return [R(0, 0, W, H, BG), T(X0, 48, 100, 20, [[(f'{n:02d}', M, MUTED)]], size=13),
            T(X0, 70, X1 - X0, 48, [[(title, B, TEXT)]], size=32, lh=1.15)]


def label(x, y, w, s, colour=MUTED):
    return T(x, y, w, 20, [[(s, M, colour)]], size=12.5)


# ---------------------------------------------------------------- content

def slide1():
    poster = BOARD / 'poster.jpg'
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), '-y', '-loglevel', 'error', '-sseof', '-1.7', '-i', str(TRAILER), '-frames:v', '1', '-q:v', '3', str(poster)], check=True)
    els = page(1, '제품 소개')
    vw = 832; vh = vw * 9 / 16
    els.append(VIDEO(TRAILER, poster, X0, TOP, vw, vh))
    rx = X0 + vw + 56; rw = X1 - rx
    els.append(label(rx, TOP, rw, '제품 개요'))
    els.append(T(rx, TOP + 28, rw, 44, [[('오행부', B, TEXT)]], size=34, lh=1.1))
    els.append(T(rx, TOP + 76, rw, 24, [[('五行符', LT, MUTED)]], size=16))
    rows = ['PC(Steam) 싱글플레이', '1인칭 한글 작도 액션 어드벤처', '29,000원', '플레이 20~30시간']
    y = TOP + 130
    for r in rows:
        els.append(HR(rx, X1, y))
        els.append(T(rx, y + 14, rw, 24, r, size=15, font=M))
        y += 54
    els.append(HR(rx, X1, y))
    return els, ('[트레일러 약 1분] 붓 한 자루를 든 도사가 되어 마법과 전설이 살아 숨 쉬는 조선을 누빕니다. 오행부는 붓으로 한글을 써서 술식을 발동하는 '
                 'PC 싱글플레이 1인칭 액션 어드벤처입니다.')


def slide2():
    els = page(2, '전통문화 활용 및 차별성')
    cols = [(SHOTS / '02_작도_나_완성.jpg', '훈민정음 해례본 제자해',
             '훈민정음의 제자 원리를 게임 시스템으로 번역했습니다. 자모를 조합하여 술식의 속성, 방향과 발현하는 방식을 정합니다.',
             '모든 술식을 외우지 않아도 반복 경험을 통해 훈민정음의 제자 원리에 따라 자연스럽게 술식을 사용합니다.'),
            (SRC / 'key302_capital.png', '한국 서예와 수묵담채',
             '가까운 산은 짙은 먹, 먼 산은 옅은 먹으로 칠해 세계를 수묵담채로 그립니다. 서예에서의 필세처럼, 마우스로 그린 획이 정확하고 빠를수록 강한 위력의 술식이 발동합니다.',
             '수묵화 속을 걸으며 마우스로 서예를 하는 듯한 경험'),
            (SRC / 'mech_buddha_2026-08.png', '한국의 역사, 사상, 종교',
             '조선의 건국이 마석의 발견으로 성사될 수 있었다는 역사적 상상력을 토대로 만든 독창적인 세계관입니다.',
             '설화 속 도깨비와 불가사리 같은 적과 싸우고, 마석으로 변한 대체역사 조선의 세계관을 체험합니다.')]
    gap = 40; cw = (X1 - X0 - 2 * gap) / 3; ih = cw * 9 / 16
    for i, (im, head, body, exp) in enumerate(cols):
        x = X0 + i * (cw + gap)
        els.append(I(im, x, TOP, cw, ih))
        els.append(T(x, TOP + ih + 20, cw, 26, [[(head, B, TEXT)]], size=17))
        els.append(T(x, TOP + ih + 54, cw, 96, body, size=13, font=LT, colour=DIM, lh=1.6))
        els.append(T(x, TOP + ih + 150, cw, 44, [[('고객 경험   ', M, MUTED), (exp, M, TEXT)]], size=13, lh=1.55))
    y = 590
    els.append(HR(X0, X1, y))
    els.append(T(X0, y + 22, 180, 24, [[('차별성', B, ACCENT)]], size=15))
    els.append(T(X0 + 180, y + 20, X1 - X0 - 180, 60,
                 [[('전통 원리를 게임 규칙으로 옮기고 아트 스타일까지 같은 원리로 묶는 ', M, TEXT), ('설계·구현 노하우', B, TEXT), ('가 차별성입니다.', M, TEXT)]], size=14.5, lh=1.6))
    return els, ('훈민정음의 제자 원리를 게임 시스템으로 번역했습니다. 자모를 조합해 술식의 속성과 방향, 발현 방식을 정하기 때문에, 모든 술식을 외우지 않아도 반복하며 자연스럽게 쓰게 됩니다. '
                 '세계는 수묵담채로 그리고, 서예의 필세처럼 마우스로 그린 획이 정확하고 빠를수록 강한 술식이 나갑니다. 세계관은 조선의 건국이 마석의 발견으로 이뤄졌다는 상상에서 출발합니다. '
                 '원전 해석부터 구현까지 한 사람이 해 왔기 때문에 90일 안에 따라 만들기 어렵습니다.')


def slide3():
    els = page(3, '목표 고객 및 구매 시나리오')
    steps = [('기존 행동', '여러 장르를 즐기고, 주말이면 스팀 신작과 할인 목록에서 혼자 할 게임을 찾습니다.'),
             ('구매 동기', '흔한 서양 판타지 게임이 아닌, 동양에서 마법을 쓰는 모습에 끌립니다.'),
             ('오행부를 고르는 이유', '스팀 페이지 트레일러에서 수묵담채의 세계 속에서 붓으로 한글을 써서 싸우는 장면을 봅니다.'),
             ('구매를 막는 요인', '새 시스템을 익혀야 한다는 부담. "글자를 모두 외워야 하나?"')]
    gap = 32; cw = (X1 - X0 - 3 * gap) / 4
    for i, (t, d) in enumerate(steps):
        x = X0 + i * (cw + gap)
        c = ACCENT if i == 3 else TEXT
        els.append(T(x, TOP, 60, 40, [[(str(i + 1), LT, MUTED if i < 3 else ACCENT)]], size=34, lh=1.0))
        els.append(T(x, TOP + 50, cw, 26, [[(t, B, c)]], size=16))
        els.append(T(x, TOP + 82, cw, 90, d, size=13, font=LT, colour=DIM, lh=1.6))
    els.append(HR(X0, X1, 330))
    lw = 340; y = 356
    els.append(label(X0, y, lw, '고객 프로필'))
    prof = [('30대 직장인 남성', '구매력이 있고 혼자 하는 게임을 선호'), ('다양한 장르', '액션과 RPG부터 인디까지'), ('스팀에서 탐색', '신작, 할인 목록과 트레일러를 보고 고름')]
    yy = y + 34
    for t, d in prof:
        els.append(T(X0, yy, lw, 22, [[(t, B, TEXT)]], size=15))
        els.append(T(X0, yy + 26, lw, 22, d, size=13, font=LT, colour=DIM))
        yy += 76
    ix = X0 + lw + 48; iw = (X1 - ix - 16) / 2
    els.append(label(ix, y, X1 - ix, '스팀 페이지 트레일러 장면'))
    els.append(I(SRC / 'key302_inn.png', ix, y + 34, iw, 250))
    els.append(I(SHOTS / '04_작도_곰.jpg', ix + iw + 16, y + 34, iw, 250))
    return els, ('목표 고객은 다양한 게임을 즐기는 30대 남성 직장인입니다. 주말 여가를 혼자 즐길 게임을 찾다가 스팀 페이지에서 오행부를 발견합니다. '
                 '흔한 서양 판타지가 아니라 동양에서 마법을 쓰는 모습에 끌리지만, 새 시스템을 익혀야 한다는 부담 때문에 구매를 망설입니다. 이 부담은 다음 장의 데모 빌드로 검증합니다.')


def slide4():
    els = page(4, 'MVP 검증 결과')
    cols = [('단계', 90), ('가설', 300), ('방법', 250), ('결과', 330), ('다음 단계', 166)]
    xs = [X0]
    for _, w in cols[:-1]:
        xs.append(xs[-1] + w)
    y = TOP
    for (name, w), x in zip(cols, xs):
        els.append(label(x, y, w - 16, name))
    rows = [('1차', '수묵담채 그래픽이 체험 의향을 만든다', '룩데브 스크린샷 설문 10명, 그래픽 선호 확인',
             [[('10명 중 8명', B, TEXT), ('이 "체험판이 나오면 해 보고 싶다"', M, TEXT)]], '그래픽과 전투 경험을 보여 줄 수 있는 플레이 시연 영상 제공'),
            ('2차', '한글 작도 전투가 매력적으로 보인다', '라이브 플레이 시연 5명, 전투 시스템 선호 확인',
             ['마석 자동차와 기계 불상에 대한 호기심', '생각보다 쉬워 보인다는 반응', [('모든 술식을 외워야 한다는 부담', B, ACCENT)]], '암기 부담을 줄인다.'),
            ('3차', '초반 술식을 제한하고 플레이하며 해금하면 암기 부담이 줄어든다', '데모 빌드(2027년 2월), 초성 5자로 시작해 강토마다 술식 해금, 시연과 같은 문항으로 비교',
             ['진행 예정'], '')]
    y += 30
    for r, (stage, hyp, how, res, nxt) in enumerate(rows):
        els.append(HR(X0, X1, y))
        rh = 104
        els.append(T(xs[0], y + 16, 80, 28, [[(stage, B, ACCENT if r == 2 else TEXT)]], size=17))
        els.append(T(xs[1], y + 16, cols[1][1] - 24, rh - 20, [[(hyp, M, TEXT)]], size=14, lh=1.55))
        els.append(T(xs[2], y + 16, cols[2][1] - 24, rh - 20, how, size=13, font=LT, colour=DIM, lh=1.55))
        els.append(T(xs[3], y + 16, cols[3][1] - 24, rh - 20, [p if isinstance(p, list) else [(p, LT, DIM)] for p in res], size=13, lh=1.5))
        els.append(T(xs[4], y + 16, cols[4][1], rh - 20, nxt, size=13, font=M, colour=TEXT if r < 2 else MUTED, lh=1.5))
        y += rh
    els.append(HR(X0, X1, y))
    iw = (X1 - X0 - 2 * 16) / 3
    for i, im in enumerate([SRC / 'key302_palace.png', STILLS / 'b1_car_250.png', SHOTS / '06_작도_뭄.jpg']):
        els.append(I(im, X0 + i * (iw + 16), y + 22, iw, 138))
    return els, ('처음에는 수묵담채 그래픽이 체험 의향을 만드는지 확인하려고 룩데브 스크린샷을 10명에게 보여 줬고, 8명이 체험판이 나오면 해 보고 싶다고 답했습니다. '
                 '다음으로 한글 작도 전투가 매력적인지 확인하려고 5명에게 라이브로 시연했습니다. 마석 자동차와 기계 불상에 대한 호기심, 생각보다 쉬워 보인다는 반응과 함께 '
                 '모든 술식을 외워야 한다는 부담이 나왔습니다. 그래서 초반에는 술식을 제한하고 플레이하며 풀어 주는 데모 빌드로 암기 부담이 줄어드는지 검증하겠습니다.')


PRICE, VAT, FEE = 29000, 1.1, .30
NET = PRICE / VAT
UNIT = NET * (1 - FEE)
MONTHS, HOURS_WEEK, WAGE = 19, 80, 12000
COSTS = [('본인 인건비', WAGE * HOURS_WEEK * MONTHS * 52 / 12, f'시급 {WAGE:,}원, 주 {HOURS_WEEK}시간, {MONTHS}개월'),
         ('AI 생산성 도구', 400000 * MONTHS, f'월 40만 원, {MONTHS}개월'), ('사운드 외주', 2800000, ''), ('일러스트 외주', 200000, ''),
         ('기타 등록비', 2000000, '플랫폼 등록과 심의')]
TOTAL = sum(c[1] for c in COSTS)
BEP = TOTAL / UNIT
TARGET, SPAN = 10000, 20000


def margin(n):
    return (n * UNIT - TOTAL) / (n * NET)


def won(v):
    man = round(v / 10000); eok, rest = divmod(man, 10000)
    return (f'{eok}억 ' if eok else '') + (f'{rest:,}만' if rest else '') + ' 원'


def slide5():
    els = page(5, '매출 구조 및 수익성')
    lw = 340
    els.append(label(X0, TOP, lw, '비용 구조'))
    els.append(T(X0, TOP - 2, lw, 22, [[(f'총 {won(TOTAL)}', B, TEXT)]], size=14, align='right'))
    yy = TOP + 36
    for name, v, d in COSTS:
        els.append(HR(X0, X0 + lw, yy - 10))
        els.append(T(X0, yy, 200, 22, [[(name, M, TEXT)]], size=14))
        els.append(T(X0 + 200, yy, lw - 200, 22, [[(won(v), M, TEXT)]], size=14, align='right'))
        if d:
            els.append(T(X0, yy + 22, lw, 20, d, size=11.5, font=LT, colour=MUTED))
        yy += 58
    els.append(HR(X0, X0 + lw, yy - 10, MUTED))
    els.append(T(X0, yy + 6, lw, 50, [[('1장당 수익  ', M, TEXT), (f'{UNIT:,.0f}원', B, TEXT)],
                                     [(f'판매가 {PRICE:,}원에서 부가세와 판매 수수료 30%를 뺀 금액', LT, MUTED)]], size=13, lh=1.6))
    cx = X0 + lw + 64; cw = X1 - cx
    els.append(label(cx, TOP, cw, '판매량별 마진율'))
    pts = [(n, round(margin(n) * 100, 1)) for n in range(5000, SPAN + 1, 1000)]
    notes = {n: (round(margin(n) * 100, 1), won(n * UNIT - TOTAL)) for n in (5000, 15000, 20000)}
    els.append(CHART('xy', cx - 10, TOP + 24, cw + 10, 340, pts=pts, mark=TARGET, marky=round(margin(TARGET) * 100, 1), labels=[5000, 10000, 15000, 20000],
                     drops=[5000, 10000, 15000, 20000], notes=notes))
    stats = [(f'{BEP:,.0f}장', '손익분기 판매량', TEXT), (f'{margin(TARGET) * 100:.1f}%', f'목표 {TARGET:,}장의 마진율', ACCENT), (won(TARGET * UNIT - TOTAL), f'목표 {TARGET:,}장의 순이익', TEXT)]
    sw = cw / 3
    for i, (v, t, c) in enumerate(stats):
        x = cx + i * sw
        els.append(T(x, 540, sw - 20, 44, [[(v, B, c)]], size=30, lh=1.1))
        els.append(T(x, 588, sw - 20, 20, t, size=12.5, font=LT, colour=MUTED))
    return els, (f'판매가 {PRICE:,}원에서 부가세와 판매 수수료 30%를 제외하면 1장당 {UNIT:,.0f}원이 남습니다. 19개월간의 인건비, AI 도구, 외주비, 등록비를 합한 총비용은 '
                 f'{won(TOTAL)}이며, 약 5,000장에서 손익분기에 도달합니다. 목표 {TARGET:,}장에서는 마진율 {margin(TARGET) * 100:.1f}%, 순이익 {won(TARGET * UNIT - TOTAL)}이고, '
                 f'20,000장까지 가면 마진율은 {margin(SPAN) * 100:.1f}%입니다.')


PLATFORMS = [('텀블벅', '8%', '플랫폼 5%와 결제 3%, VAT 별도', '창작자 크라우드펀딩 플랫폼. 데모와 함께 선판매와 후원을 받습니다.', '전통문화와 인디 콘텐츠 후원자', 20, '8E8A84'),
             ('스팀', '30%', '매출 1천만 달러 이하 구간, 등록비 100달러', 'PC 게임 유통 플랫폼. 위시리스트와 넥스트 페스트로 알립니다.', '전 세계 액션, 인디 게이머', 70, 'F2F0EC'),
             ('스토브 인디', '15%', '스마일게이트 발표(2023) 기준', '스마일게이트의 국내 인디 게임 플랫폼. 기획전과 커뮤니티가 있습니다.', '국내 인디 게임 이용자', 10, '4A4845')]


def slide6():
    els = page(6, '사업화 계획 및 판매 채널')
    miles = [('2027년 2월', '데모 공개와 펀딩', '텀블벅 펀딩, 스팀 데모', '텀블벅, 스팀'),
             ('2027년 8월', '얼리 액세스', '선판매와 피드백 반영', '스팀, 스토브 인디'),
             ('2028년 2월', '정식 출시', '5개 강토, 3개 결말', '스팀, 스토브 인디')]
    gap = 40; cw = (X1 - X0 - 2 * gap) / 3
    for i, (d, t, sub, tgt) in enumerate(miles):
        x = X0 + i * (cw + gap)
        els.append(T(x, TOP, cw, 20, [[(d, M, ACCENT if i == 0 else MUTED)]], size=13))
        els.append(T(x, TOP + 26, cw, 30, [[(t, B, TEXT)]], size=20))
        els.append(T(x, TOP + 64, cw, 22, sub, size=13, font=LT, colour=DIM))
        els.append(T(x, TOP + 92, cw, 22, [[('타겟 플랫폼   ', M, MUTED), (tgt, M, TEXT)]], size=13))
    els.append(HR(X0, X1, 290))
    pw = 250; cw2 = (X1 - X0 - pw - 3 * gap) / 3; y2 = 316
    for i, (name, fee, note, trait, cust, share, c) in enumerate(PLATFORMS):
        x = X0 + i * (cw2 + gap)
        els.append(T(x, y2, cw2, 26, [[(name, B, TEXT)]], size=18))
        els.append(label(x, y2 + 40, cw2, '예상 수수료'))
        els.append(T(x, y2 + 60, cw2, 44, [[(fee, B, TEXT)]], size=32, lh=1.0))
        els.append(T(x, y2 + 104, cw2, 20, note, size=11, font=LT, colour=MUTED))
        els.append(T(x, y2 + 140, cw2, 70, trait, size=13, font=LT, colour=DIM, lh=1.6))
        els.append(T(x, y2 + 222, cw2, 44, [[('고객   ', M, MUTED), (cust, M, TEXT)]], size=13, lh=1.5))
    px = X1 - pw
    els.append(label(px, y2, pw, '예상 매출 비율'))
    els.append(CHART('pie', px + 35, y2 + 36, 180, 180, cats=[p[0] for p in PLATFORMS], vals=[p[5] for p in PLATFORMS], colours=[p[6] for p in PLATFORMS]))
    yy = y2 + 236
    for p in sorted(PLATFORMS, key=lambda q: -q[5]):
        els.append(R(px, yy + 6, 10, 10, p[6]))
        els.append(T(px + 18, yy, 150, 22, p[0], size=13, font=M))
        els.append(T(px + 150, yy, pw - 150, 22, [[(f'{p[5]}%', B, TEXT)]], size=13, align='right'))
        yy += 26
    return els, ('2027년 2월에는 텀블벅 펀딩과 스팀 데모로 초기 수요를 확인합니다. 8월에는 스팀과 스토브 인디에서 얼리 액세스로 판매를 시작하고, 2028년 2월 두 플랫폼에서 정식 출시합니다. '
                 '매출의 70%는 스팀, 20%는 텀블벅, 10%는 스토브 인디에서 나올 것으로 봅니다.')


def slide7():
    els = page(7, '팀 역량 및 보완 계획')
    gap = 64; cw = (X1 - X0 - gap) / 2
    groups = [('대표 한예준', '기획, 개발, 아트', [('유학동양학 전공', '동양 철학의 현대적 의미에 대해 항상 고민해 왔습니다.'),
                                             ('게임 개발 부트캠프 수료', 'Unity/C#을 이용한 게임 개발 및 협업 경험 다수'),
                                             ('외주 개발', '수료 후 외주 개발로 실무 경험을 쌓았습니다.'),
                                             ('오행부 개발 현황', '5개 강토 월드, 한글 작도 인식, 전투와 요괴, 마석 자동차, 수묵 렌더링까지 플레이할 수 있는 빌드를 만들었습니다.')]),
              ('보완 계획', '2027년 1분기까지', [('사운드', 'AI로 시안을 만든 뒤 외주로 완성합니다.'),
                                           ('마케팅', '창업 프로그램의 멘토링과 펀딩, 게임 행사와 전시로 노출과 피드백을 얻습니다.')])]
    for i, (head, sub, items) in enumerate(groups):
        x = X0 + i * (cw + gap)
        els.append(T(x, TOP, cw, 30, [[(head, B, TEXT), ('   ' + sub, LT, MUTED)]], size=20))
        yy = TOP + 58
        for t, d in items:
            els.append(HR(x, x + cw, yy - 16))
            els.append(T(x, yy, cw, 26, [[(t, B, TEXT)]], size=16))
            if d:
                els.append(T(x, yy + 30, cw, 50, d, size=13.5, font=LT, colour=DIM, lh=1.6))
            yy += 104 if d else 64
    return els, ('저는 유학동양학을 전공하며 동양 철학의 현대적 의미를 늘 고민해 왔고, 게임 개발 부트캠프에서 Unity와 C#으로 여러 차례 게임을 개발하고 협업했으며, 수료 후 외주 개발을 진행했습니다. '
                 '현재 오행부는 다섯 강토와 작도 인식, 전투, 수묵 렌더링까지 플레이할 수 있습니다. 부족한 사운드는 AI로 시안을 만든 뒤 외주로 완성하고, '
                 '마케팅은 창업 프로그램과 게임 행사를 활용해 2027년 1분기 안에 채우겠습니다.')


SLIDES = [slide1, slide2, slide3, slide4, slide5, slide6, slide7]


# ---------------------------------------------------------------- HTML board

def paras(runs, font, colour):
    if isinstance(runs, str):
        runs = [runs]
    out = []
    for p in runs:
        if isinstance(p, str):
            p = [(p, font, colour)]
        out.append(''.join(f'<span style="font-family:\'{f}\';color:#{c}">{html.escape(t)}</span>' for t, f, c in p))
    return out


def svg_chart(e):
    w, h = e['w'], e['h']
    if e['kind'] == 'xy':
        pts = e['pts']; l, r, t, b = 48, 20, 26, 34
        x0, x1 = pts[0][0], pts[-1][0]
        sx = lambda v: l + (v - x0) / (x1 - x0) * (w - l - r)
        sy = lambda v: t + (1 - v / 60) * (h - t - b)
        g = ''.join(f'<line x1="{l}" x2="{w - r}" y1="{sy(v):.1f}" y2="{sy(v):.1f}" stroke="#{RULE}"/><text x="{l - 10}" y="{sy(v) + 4:.1f}" text-anchor="end">{v}%</text>' for v in range(0, 70, 10))
        line = ' '.join(f'{sx(a):.1f},{sy(v):.1f}' for a, v in pts)
        mx, my = sx(e['mark']), sy(e['marky'])
        dots = ''.join(f'<circle cx="{sx(a):.1f}" cy="{sy(v):.1f}" r="3" fill="#{TEXT}"/>' for a, v in pts)
        val = dict(pts)
        drops = ''.join(f'<line x1="{sx(a):.1f}" x2="{sx(a):.1f}" y1="{sy(val[a]):.1f}" y2="{sy(0):.1f}" stroke="#5A5752" stroke-dasharray="4,4" stroke-width="1"/>' for a in e['drops'] if a != e['mark'])
        lab = f'<text x="{mx:.1f}" y="{my - 12:.1f}" text-anchor="middle" style="fill:#{ACCENT};font-weight:700">{e["marky"]}%</text>'
        for a, (m_, p_) in e['notes'].items():
            anchor = 'end' if a == pts[-1][0] else 'start' if a == pts[0][0] else 'middle'
            lab += (f'<text x="{sx(a):.1f}" y="{sy(val[a]) - 24:.1f}" text-anchor="{anchor}" style="font-size:11px">{m_}%</text>'
                    f'<text x="{sx(a):.1f}" y="{sy(val[a]) - 10:.1f}" text-anchor="{anchor}" style="font-size:11px">{p_}</text>')
        xl = ''.join(f'<text x="{sx(a):.1f}" y="{h - 10}" text-anchor="middle">{a:,}장</text>' for a in e['labels'])
        return (f'<svg width="{w}" height="{h}" style="font:12px \'{LT}\';fill:#{MUTED}">{g}{drops}<polyline points="{line}" fill="none" stroke="#{TEXT}" stroke-width="2"/>{dots}'
                f'<line x1="{mx:.1f}" x2="{mx:.1f}" y1="{my:.1f}" y2="{sy(0):.1f}" stroke="#{ACCENT}" stroke-dasharray="5,4" stroke-width="1.5"/>'
                f'<circle cx="{mx:.1f}" cy="{my:.1f}" r="6" fill="#{ACCENT}"/>{lab}{xl}</svg>')
    tot = sum(e['vals']); a0 = -math.pi / 2; c = w / 2; out = ''
    for v, col in zip(e['vals'], e['colours']):
        a1 = a0 + 2 * math.pi * v / tot; large = 1 if a1 - a0 > math.pi else 0
        out += f'<path d="M{c},{c} L{c + c * math.cos(a0):.1f},{c + c * math.sin(a0):.1f} A{c},{c} 0 {large} 1 {c + c * math.cos(a1):.1f},{c + c * math.sin(a1):.1f} Z" fill="#{col}" stroke="#{BG}" stroke-width="2"/>'
        a0 = a1
    return f'<svg width="{w}" height="{h}">{out}</svg>'


def render_html(pages):
    faces = ''.join(f"@font-face{{font-family:'{n}';src:url('file:///C:/Windows/Fonts/{f}.ttf')}}" for n, f in FILES.items())
    out = [f'<!doctype html><html><head><meta charset="utf-8"><title>오행부 피칭 v6</title><style>{faces}'
           'body{margin:0;background:#3a3a3a;display:flex;flex-direction:column;align-items:center;gap:24px;padding:24px}'
           '.s{position:relative;width:1280px;height:720px;overflow:hidden}.s>*{position:absolute;box-sizing:border-box;margin:0}p{margin:0}</style></head><body>']
    for els, _ in pages:
        out.append('<div class="s">')
        for e in els:
            st = f"left:{e['x']:.1f}px;top:{e['y']:.1f}px;width:{e['w']:.1f}px;height:{e['h']:.1f}px"
            if e['k'] == 'rect':
                out.append(f'<div style="{st};background:#{e["fill"]}"></div>')
            elif e['k'] == 'text':
                ps = ''.join(f'<p>{p}</p>' for p in paras(e['runs'], e['font'], e['colour']))
                out.append(f'<div style="{st};font-size:{e["size"]}px;line-height:{e["lh"]};text-align:{e["align"]}">{ps}</div>')
            elif e['k'] in ('img', 'video'):
                src = e['poster'] if e['k'] == 'video' else e['path']
                out.append(f'<img src="{src.as_uri()}" style="{st};object-fit:cover">')
            elif e['k'] == 'chart':
                out.append(f'<div style="{st}">{svg_chart(e)}</div>')
        out.append('</div>')
    out.append('</body></html>')
    path = BOARD / 'board.html'; path.write_text(''.join(out), encoding='utf-8')
    return path


# ---------------------------------------------------------------- PPTX

def px(v):
    return Emu(int(round(v * 9525)))


def rgb(h):
    return RGBColor.from_string(h)


def run_font(r, font, size, colour):
    r.font.size = Pt(size * .75); r.font.color.rgb = rgb(colour); r.font.name = font
    rpr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {}); rpr.append(el)
        el.set('typeface', font)


def axis_font(ax, size, colour):
    ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.color.rgb = rgb(colour); ax.tick_labels.font.name = LT


def xy_chart(s, e):
    data = XyChartData()
    curve = data.add_series('마진율')
    for a, v in e['pts']:
        curve.add_data_point(a, v)
    val = dict(e['pts'])
    for a in [e['mark']] + [a for a in e['drops'] if a != e['mark']]:
        d_ = data.add_series(f'{a:,}장'); d_.add_data_point(a, val[a]); d_.add_data_point(a, 0)
    ch = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, px(e['x']), px(e['y']), px(e['w']), px(e['h']), data).chart
    ch.has_title = False; ch.has_legend = False; ch.font.name = LT; ch.font.color.rgb = rgb(MUTED)
    series = list(ch.plots[0].series)
    c = series[0]
    c.format.line.color.rgb = rgb(TEXT); c.format.line.width = Pt(2); c.smooth = False
    c.marker.style = XL_MARKER_STYLE.CIRCLE; c.marker.size = 5
    c.marker.format.fill.solid(); c.marker.format.fill.fore_color.rgb = rgb(TEXT); c.marker.format.line.color.rgb = rgb(TEXT)
    for j, d in enumerate(series[1:]):
        first = j == 0                              # the target line keeps the accent, the others stay quiet
        d.format.line.color.rgb = rgb(ACCENT if first else '5A5752'); d.format.line.width = Pt(1.5 if first else .75)
        d.format.line.dash_style = MSO_LINE.DASH; d.smooth = False; d.marker.style = XL_MARKER_STYLE.NONE
    for i, (a, v) in enumerate(e['pts']):
        if a == e['mark']:
            dl = c.points[i].data_label
            dl.has_text_frame = True; dl.text_frame.text = f'{v}%'; dl.position = XL_LABEL_POSITION.ABOVE
            r = dl.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = rgb(ACCENT)
        elif a in e['notes']:
            m_, p_ = e['notes'][a]
            dl = c.points[i].data_label
            dl.has_text_frame = True; dl.text_frame.text = f'{m_}%\n{p_}'
            dl.position = XL_LABEL_POSITION.RIGHT if a == e['pts'][0][0] else XL_LABEL_POSITION.ABOVE
            for para in dl.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(8); r.font.bold = False; r.font.color.rgb = rgb(MUTED)
        if a == e['mark']:
            p = c.points[i]; p.marker.size = 10; p.marker.format.fill.solid(); p.marker.format.fill.fore_color.rgb = rgb(ACCENT); p.marker.format.line.color.rgb = rgb(ACCENT)
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = e['pts'][0][0], e['pts'][-1][0], 5000
    xa.tick_labels.number_format = '#,##0"장"'; xa.tick_labels.number_format_is_linked = False
    xa.format.line.color.rgb = rgb(RULE); xa.has_major_gridlines = False; axis_font(xa, 9, MUTED)
    ya.minimum_scale, ya.maximum_scale, ya.major_unit = 0, 60, 10
    ya.tick_labels.number_format = '0"%"'; ya.tick_labels.number_format_is_linked = False
    ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = rgb(RULE); ya.major_gridlines.format.line.width = Pt(.5)
    ya.format.line.fill.background(); axis_font(ya, 9, MUTED)


def render_pptx(pages):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    for els, note in pages:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for e in els:
            if e['k'] == 'rect':
                sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e['fill']); sh.line.fill.background(); sh.shadow.inherit = False
            elif e['k'] == 'text':
                tb = s.shapes.add_textbox(px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                runs = e['runs'] if not isinstance(e['runs'], str) else [e['runs']]
                for i, p in enumerate(runs):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = {'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}[e['align']]
                    para.line_spacing = e['lh'] * .83
                    for t, f, c in ([(p, e['font'], e['colour'])] if isinstance(p, str) else p):
                        r = para.add_run(); r.text = t; run_font(r, f, e['size'], c)
            elif e['k'] == 'img':
                iw, ih = Image.open(e['path']).size
                pic = s.shapes.add_picture(str(e['path']), px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                if iw / ih > e['w'] / e['h']:
                    c = (1 - (e['w'] / e['h']) / (iw / ih)) / 2; pic.crop_left = pic.crop_right = c
                else:
                    c = (1 - (iw / ih) / (e['w'] / e['h'])) / 2; pic.crop_top = pic.crop_bottom = c
            elif e['k'] == 'video':
                s.shapes.add_movie(str(e['path']), px(e['x']), px(e['y']), px(e['w']), px(e['h']), poster_frame_image=str(e['poster']), mime_type='video/mp4')
            elif e['k'] == 'chart':
                if e['kind'] == 'xy':
                    xy_chart(s, e)
                else:
                    data = CategoryChartData(); data.categories = e['cats']; data.add_series('비율', [v / 100 for v in e['vals']])
                    ch = s.shapes.add_chart(XL_CHART_TYPE.PIE, px(e['x']), px(e['y']), px(e['w']), px(e['h']), data).chart
                    ch.has_title = False; ch.has_legend = False
                    plot = ch.plots[0]; plot.has_data_labels = False
                    for i, col in enumerate(e['colours']):
                        pt = plot.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
                        pt.format.line.color.rgb = rgb(BG); pt.format.line.width = Pt(1.5)
        if '--no-notes' not in sys.argv: s.notes_slide.notes_text_frame.text = note   # --no-notes: the deck without the talk script
    prs.save(DECK)
    if '--no-embed' not in sys.argv:
        ps = (f'$app = New-Object -ComObject PowerPoint.Application; $p = $app.Presentations.Open("{DECK}", $false, $false, $false); '
              f'$p.SaveAs("{DECK}", 24, -1); $p.Close()')
        subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True)
    return DECK


def main():
    BOARD.mkdir(parents=True, exist_ok=True)
    pages = [f() for f in SLIDES]
    print(render_html(pages))
    if '--html-only' not in sys.argv:
        print(render_pptx(pages))
    print(f'margin 10k {margin(TARGET) * 100:.1f}%  20k {margin(SPAN) * 100:.1f}%  bep {BEP:,.0f}')


if __name__ == '__main__':
    main()
