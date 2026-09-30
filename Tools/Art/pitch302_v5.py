"""#302 pitch deck v5: one layout spec rendered twice, as an HTML design board and as the editable .pptx.

Design (frontend-design skill, two passes): the subject's own material is the woodblock page of 훈민정음 해례본.
The left band is the 판심 (fold column): 어미 marks top and bottom, the section name written vertically, the page
numeral. Content columns are divided only by 계선 (thin vertical rules). No cards, shadows, arrows or dot-strings.
Palette 장지 EDEBE6 · 먹 1C1B1A · 담묵 85817A · 계선 BDB7AC · 인주 9E3027 (one or two uses per page).
Type KoPub 바탕 Bold for headings, KoPub 돋움 for text. Canvas 1280 x 720 px = 13.333 x 7.5 in (96 px/in).

usage: python Tools/Art/pitch302_v5.py [--trailer path.mp4] [--no-embed]
Output: Art/Presentation302/v5/board.html, Art/Presentation302/오행부_피칭_v5.pptx
"""
import html
import subprocess
import sys
from pathlib import Path

import imageio_ffmpeg
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'Art/Presentation302'
BOARD = OUT / 'v5'
SRC = OUT / 'src'
SHOTS = ROOT / 'Art/Presentation297/Screenshots'
STILLS = ROOT / 'Art/Presentation297/Stills'
TRAILER = Path(sys.argv[sys.argv.index('--trailer') + 1]) if '--trailer' in sys.argv else ROOT / 'Art/Presentation297/Trailer/오행부_트레일러_1분_드라이브용.mp4'
DECK = OUT / '오행부_피칭_v5.pptx'

PAPER, INK, GREY, RULE, SEAL = 'EDEBE6', '1C1B1A', '85817A', 'BDB7AC', '9E3027'
HEAD, HEAD_M, BODY, BODY_B = 'KoPubBatang Bold', 'KoPubBatang Medium', 'KoPubDotum Medium', 'KoPubDotum Bold'
W, H = 1280, 720
X0, X1, Y0 = 150, 1236, 52          # content left/right edge, heading top


# ---------------------------------------------------------------- the layout model

def T(x, y, w, h, runs, size=15, font=BODY, colour=INK, align='left', lh=1.45, anchor='top'):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, font, colour)."""
    return dict(k='text', x=x, y=y, w=w, h=h, runs=runs, size=size, font=font, colour=colour, align=align, lh=lh, anchor=anchor)


def Rl(x, y, w, h, fill):
    return dict(k='rect', x=x, y=y, w=w, h=h, fill=fill)


def VR(x, y0, y1, colour=RULE):          # 계선
    return Rl(x, y0, 1, y1 - y0, colour)


def HR(x0, x1, y, colour=RULE):
    return Rl(x0, y, x1 - x0, 1, colour)


def I(path, x, y, w, h):
    return dict(k='img', path=Path(path), x=x, y=y, w=w, h=h)


def POLY(points, fill):
    return dict(k='poly', points=points, fill=fill)


def CHART(kind, x, y, w, h, **data):
    return dict(k='chart', kind=kind, x=x, y=y, w=w, h=h, **data)


def VIDEO(path, poster, x, y, w, h):
    return dict(k='video', path=Path(path), poster=Path(poster), x=x, y=y, w=w, h=h)


NUM = '一二三四五六七'


def eomi(cx, y, down=True, w=44, h=22):
    """어미: a black block ending in a fish-tail notch (pointing down at the top of the 판심, up at the bottom)."""
    l, r = cx - w / 2, cx + w / 2
    if down:
        pts = [(l, y), (r, y), (r, y + h), (cx + w * .18, y + h), (cx, y + h * .55), (cx - w * .18, y + h), (l, y + h)]
    else:
        pts = [(l, y + h), (r, y + h), (r, y), (cx + w * .18, y), (cx, y + h * .45), (cx - w * .18, y), (l, y)]
    return POLY(pts, INK)


def frame(n, section, heading):
    """판심 band + heading. Returns the element list of the page frame."""
    cx = 58
    els = [Rl(0, 0, W, H, PAPER), VR(22, 28, 692, INK), VR(94, 28, 692, INK), VR(24, 28, 692, RULE), VR(92, 28, 692, RULE),
           eomi(cx, 52), eomi(cx, 646, down=False)]
    els.append(T(cx - 16, 92, 32, 84, [[(c, HEAD_M, GREY)] for c in '五行符'], size=17, align='center', lh=1.12))
    els.append(HR(cx - 8, cx + 8, 184, INK))
    sec = list(section.replace(' ', ''))
    els.append(T(cx - 16, 196, 32, 26 * len(sec) + 10, [[(c, HEAD, INK)] for c in sec], size=17, align='center', lh=1.12))
    els.append(T(cx - 16, 598, 32, 34, [[(NUM[n - 1], HEAD, SEAL)]], size=22, align='center', lh=1.0))
    els.append(T(X0, Y0, X1 - X0, 48, [[(heading, HEAD, INK)]], size=32, lh=1.15))
    return els


# ---------------------------------------------------------------- content

def slide1():
    poster = BOARD / 'poster.jpg'
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), '-y', '-loglevel', 'error', '-ss', '52', '-i', str(TRAILER), '-frames:v', '1', '-q:v', '3', str(poster)], check=True)
    els = frame(1, '제품소개', '제품 소개')
    vw = 846; vh = vw * 9 / 16; vy = 124
    els.append(VIDEO(TRAILER, poster, X0, vy, vw, vh))
    rx = X0 + vw + 34; els.append(VR(X0 + vw + 17, vy, vy + vh))
    blocks = [('제품명', [[('오행부', HEAD, INK)], [('五行符', HEAD_M, GREY)]], 30),
              ('제품 형태', ['PC(Steam) 싱글플레이', '1인칭 한글 작도 액션', '29,000원', '플레이 20~30시간'], 15),
              ('핵심 경험', ['붓 한 자루를 든 도사가 되어', '마법과 전설이 살아 숨 쉬는', '조선을 누빈다'], 15)]
    y = vy
    for label, runs, size in blocks:
        els.append(T(rx, y, X1 - rx, 20, label, size=12, colour=GREY))
        hh = 86 if size > 20 else (len(runs) * 23 + 4)
        els.append(T(rx, y + 22, X1 - rx, hh, runs, size=size, font=BODY if size < 20 else HEAD, lh=1.35 if size < 20 else 1.1))
        y += 22 + hh + 26
        if label != '핵심 경험':
            els.append(HR(rx, X1, y - 14))
    return els, '[트레일러 약 1분] 오행부는 붓으로 한글을 써서 술식을 발동하는 PC 싱글플레이 액션 어드벤처입니다. 플레이어는 붓 한 자루를 든 도사가 되어 마법과 전설이 살아 숨 쉬는 조선을 누빕니다.'


def slide2():
    els = frame(2, '전통문화', '전통문화 활용 및 차별성')
    cols = [(SHOTS / '02_작도_나_완성.jpg', '훈민정음 해례본 제자해',
             '제자해의 오행 배정에 따라 ㄱ, ㄴ, ㅁ, ㅅ, ㅇ은 목, 화, 토, 금, 수 속성을 갖습니다. 자모를 조합하는 규칙이 그대로 술식 문법입니다.',
             '자모를 조합해 새 술식을 찾아냅니다.'),
            (SRC / 'key302_capital.png', '한국 서예와 수묵담채',
             '붓으로 쓴 획이 정확하고 빠를수록 위력이 셉니다. 가까운 산은 짙은 먹, 먼 산은 옅은 먹으로 칠해 세계를 수묵담채로 그립니다.',
             '수묵화 속을 걸으며 직접 붓으로 씁니다.'),
            (SRC / 'mech_buddha_2026-08.png', '역사, 사상, 종교',
             '마석으로 산업화한 조선이 배경이고, 오행의 상생과 상극이 전투 규칙입니다. 성황당, 기계 불상, 도깨비, 불가사리가 세계와 적으로 나옵니다.',
             '설화 속 도깨비와 불가사리를 적으로 만납니다.')]
    gap = 36; cw = (X1 - X0 - 2 * gap) / 3; y = 124
    for i, (im, head, body, exp) in enumerate(cols):
        x = X0 + i * (cw + gap)
        if i:
            els.append(VR(x - gap / 2, y, 560))
        els.append(I(im, x, y, cw, cw * 9 / 16))
        iy = y + cw * 9 / 16
        els.append(T(x, iy + 18, cw, 28, [[(head, HEAD, INK)]], size=19))
        els.append(T(x, iy + 54, cw, 96, body, size=13.5, lh=1.55))
        els.append(T(x, iy + 158, cw, 44, [[('고객 경험  ', BODY_B, SEAL if i == 0 else INK), (exp, BODY, INK)]], size=13.5))
    els.append(HR(X0, X1, 584, INK))
    els.append(T(X0, 602, 120, 24, [[('차별성', HEAD, INK)]], size=17))
    els.append(T(X0 + 130, 602, X1 - X0 - 130, 60,
                 '전통 원리를 게임 규칙으로 옮기고 아트 스타일까지 같은 원리로 묶는 설계·구현 노하우가 차별성입니다. 원전 해석부터 구현까지 한 사람이 해 왔기 때문에 90일 안에 따라 만들기 어렵습니다.',
                 size=14.5, lh=1.55))
    return els, ('훈민정음 해례본 제자해는 소리 나는 자리에 따라 자음에 오행을 배정했습니다. 오행부는 이 배정을 그대로 술식 문법으로 씁니다. 붓글씨의 필세는 술식의 위력이 되고, '
                 '수묵담채는 화면을 그리는 방식이 되며, 역사와 민간신앙은 세계와 적이 됩니다. 원전 해석부터 구현까지 한 사람이 해 왔기 때문에 90일 안에 따라 만들기 어렵습니다.')


def slide3():
    els = frame(3, '목표고객', '목표 고객 및 구매 시나리오')
    steps = [('기존 행동', '여러 장르를 즐기고, 주말이면 스팀 신작과 할인 목록에서 혼자 할 게임을 찾습니다.'),
             ('구매 동기', '서양 판타지가 아닌 세계와 처음 보는 조작 방식을 찾습니다.'),
             ('오행부를 고르는 이유', '스팀 페이지 트레일러에서 붓으로 한글을 써서 싸우는 장면을 봅니다.'),
             ('구매를 막는 요인', '새 시스템을 익혀야 한다는 부담. "글자를 모두 외워야 하나?"')]
    gap = 30; cw = (X1 - X0 - 3 * gap) / 4; y = 124
    for i, (t, d) in enumerate(steps):
        x = X0 + i * (cw + gap)
        if i:
            els.append(VR(x - gap / 2, y, 300))
        c = SEAL if i == 3 else INK
        els.append(T(x, y, 40, 36, [[(str(i + 1), HEAD, c)]], size=28, lh=1.0))
        els.append(T(x, y + 44, cw, 26, [[(t, HEAD, c)]], size=17))
        els.append(T(x, y + 78, cw, 100, d, size=13.5, lh=1.55))
    els.append(HR(X0, X1, 330, INK))
    lw = 330
    els.append(T(X0, 350, lw, 26, [[('고객 프로필', HEAD, INK)]], size=17))
    prof = [('30대 직장인 남성', '구매력이 있고 혼자 하는 게임을 선호'), ('다양한 장르', '액션과 RPG부터 인디까지'), ('스팀에서 탐색', '신작, 할인 목록과 트레일러를 보고 고름')]
    yy = 392
    for t, d in prof:
        els.append(T(X0, yy, lw, 22, [[(t, BODY_B, INK)]], size=14))
        els.append(T(X0, yy + 24, lw, 22, d, size=13, colour=GREY))
        yy += 70
    els.append(VR(X0 + lw + 18, 350, 660))
    ix = X0 + lw + 36; iw = (X1 - ix - 16) / 2; ih = 236
    els.append(T(ix, 350, X1 - ix, 26, [[('스팀 페이지 트레일러 장면', HEAD, INK)]], size=17))
    els.append(I(SRC / 'key302_inn.png', ix, 392, iw, ih))
    els.append(I(SHOTS / '04_작도_곰.jpg', ix + iw + 16, 392, iw, ih))
    return els, ('목표 고객은 다양한 게임을 즐기는 30대 남성 직장인입니다. 주말 여가를 혼자 즐길 게임을 찾다가 스팀 페이지에서 오행부를 발견합니다. '
                 '처음 보는 세계와 조작 방식에 끌리지만, 새 시스템을 익혀야 한다는 부담 때문에 구매를 망설입니다. 이 부담은 다음 장의 데모 빌드로 검증합니다.')


def slide4():
    els = frame(4, '검증결과', 'MVP 검증 결과')
    gap = 36; cw = (X1 - X0 - 2 * gap) / 3; y = 124
    for i in range(1, 3):
        els.append(VR(X0 + i * (cw + gap) - gap / 2, y, 660))
    x = X0
    els.append(T(x, y, cw, 20, '스크린샷 설문, 10명', size=12, colour=GREY))
    els.append(T(x, y + 22, cw, 80, [[('8', HEAD, INK), (' / 10', HEAD_M, GREY)]], size=60, lh=1.05))
    els.append(T(x, y + 112, cw, 70, '룩데브 스크린샷을 본 10명 중 8명이 "체험판이 나오면 해 보고 싶다"고 답했습니다.', size=14, lh=1.55))
    x = X0 + cw + gap
    els.append(T(x, y, cw, 20, '라이브 시연, 5명', size=12, colour=GREY))
    q = [('호기심', '마석 자동차와 기계 불상이 궁금하다', INK), ('난이도', '생각보다 쉬워 보인다', INK), ('부담', '술식을 모두 외워야 하는지 걱정된다', SEAL)]
    yy = y + 30
    for t, d, c in q:
        els.append(T(x, yy, cw, 22, [[(t, BODY_B, c)]], size=13))
        els.append(T(x, yy + 22, cw, 30, [[(d, HEAD_M, INK)]], size=16.5))
        yy += 58
    x = X0 + 2 * (cw + gap)
    els.append(T(x, y, cw, 20, '데모 빌드 검증, 2027년 2월', size=12, colour=GREY))
    plan = [('가설', '처음에 술식을 제한하고 플레이하며 풀어 주면 암기 부담이 줄어든다'), ('해금', '초성 5자로 시작해 강토를 깰 때마다 종성을 얻고 술식 10종이 한 번에 열린다'),
            ('측정', '시연 때와 같은 질문으로 암기 부담을 비교한다')]
    yy = y + 30
    for t, d in plan:
        els.append(T(x, yy, cw, 22, [[(t, BODY_B, INK)]], size=13))
        els.append(T(x, yy + 22, cw, 44, d, size=13.5, lh=1.5))
        yy += 64
    for i, im in enumerate([SRC / 'key302_palace.png', STILLS / 'b1_car_250.png', SHOTS / '06_작도_뭄.jpg']):
        els.append(I(im, X0 + i * (cw + gap), 446, cw, 214))
    return els, ('룩데브 스크린샷을 본 10명 중 8명이 체험판이 나오면 해 보고 싶다고 답했습니다. 5명에게 라이브로 시연하자 마석 자동차와 기계 불상이 궁금하다는 반응, '
                 '생각보다 쉬워 보인다는 반응이 나왔고, 술식을 모두 외워야 하느냐는 걱정도 나왔습니다. 그래서 처음에는 술식을 제한하고 플레이하며 풀어 주는 데모 빌드를 만들어, '
                 '암기 부담이 실제로 줄어드는지 확인하겠습니다.')


PRICE, VAT, FEE = 29000, 1.1, .30
NET = PRICE / VAT
UNIT = NET * (1 - FEE)
MONTHS, HOURS_WEEK, WAGE = 19, 80, 12000
COSTS = [('본인 인건비', WAGE * HOURS_WEEK * MONTHS * 52 / 12, f'시급 {WAGE:,}원, 주 {HOURS_WEEK}시간, {MONTHS}개월'),
         ('AI 생산성 도구', 400000 * MONTHS, f'월 40만 원, {MONTHS}개월'), ('사운드 외주', 2800000, ''), ('일러스트 외주', 200000, ''),
         ('기타 등록비', 2000000, '플랫폼 등록과 심의')]
TOTAL = sum(c[1] for c in COSTS)
BEP = TOTAL / UNIT
TARGET = 10000


def margin(n):
    return (n * UNIT - TOTAL) / (n * NET)


def won(v):
    man = round(v / 10000); eok, rest = divmod(man, 10000)
    return (f'{eok}억 ' if eok else '') + (f'{rest:,}만' if rest else '') + ' 원'


def slide5():
    els = frame(5, '수익성', '매출 구조 및 수익성')
    lw = 360; y = 124
    els.append(T(X0, y, lw, 26, [[('비용 구조', HEAD, INK)]], size=17))
    els.append(T(X0, y, lw, 26, [[(f'총 {won(TOTAL)}', BODY_B, INK)]], size=14, align='right'))
    yy = y + 44
    for name, v, d in COSTS:
        els.append(HR(X0, X0 + lw, yy - 8))
        els.append(T(X0, yy, 200, 22, [[(name, BODY_B, INK)]], size=14))
        els.append(T(X0 + 200, yy, lw - 200, 22, won(v), size=14, align='right'))
        if d:
            els.append(T(X0, yy + 22, lw, 20, d, size=11.5, colour=GREY))
        yy += 58
    els.append(HR(X0, X0 + lw, yy - 8, INK))
    els.append(T(X0, yy + 4, lw, 44, [[('1장당 수익  ', BODY_B, INK), (f'{UNIT:,.0f}원', HEAD, INK)],
                                     [(f'판매가 {PRICE:,}원에서 부가세와 판매 수수료 30%를 뺀 금액', BODY, GREY)]], size=13, lh=1.5))
    els.append(VR(X0 + lw + 24, y, 660))
    cx = X0 + lw + 48; cw = X1 - cx
    els.append(T(cx, y, cw, 26, [[('판매량별 마진율', HEAD, INK)]], size=17))
    counts = list(range(5000, TARGET + 1, 1000))
    els.append(CHART('line', cx - 8, y + 36, cw + 8, 320, cats=[f'{n:,}장' for n in counts], vals=[round(margin(n) * 100, 1) for n in counts]))
    stats = [(f'{BEP:,.0f}장', '손익분기 판매량'), (f'{margin(TARGET) * 100:.1f}%', f'목표 {TARGET:,}장의 마진율'), (won(TARGET * UNIT - TOTAL), f'목표 {TARGET:,}장의 순이익')]
    sw = cw / 3
    for i, (v, t) in enumerate(stats):
        x = cx + i * sw
        if i:
            els.append(VR(x - 12, 496, 660))
        els.append(T(x, 506, sw - 24, 44, [[(v, HEAD, SEAL if i == 1 else INK)]], size=30, lh=1.1))
        els.append(T(x, 556, sw - 24, 20, t, size=12.5, colour=GREY))
    return els, (f'판매가 {PRICE:,}원에서 부가세와 판매 수수료 30%를 제외하면 1장당 {UNIT:,.0f}원이 남습니다. 19개월간의 인건비, AI 도구, 외주비, 등록비를 합한 총비용은 '
                 f'{won(TOTAL)}이며, 약 5,000장에서 손익분기에 도달합니다. 이후 판매량이 1,000장 늘어날 때마다 마진율이 오르며, 목표 {TARGET:,}장에서는 '
                 f'마진율 {margin(TARGET) * 100:.1f}%, 순이익 {won(TARGET * UNIT - TOTAL)}입니다.')


PLATFORMS = [('텀블벅', '8%', '플랫폼 5%와 결제 3%, VAT 별도', '창작자 크라우드펀딩 플랫폼. 데모와 함께 선판매와 후원을 받습니다.', '전통문화와 인디 콘텐츠 후원자', 20, '8A7A55'),
             ('스팀', '30%', '매출 1천만 달러 이하 구간, 등록비 100달러', 'PC 게임 유통 플랫폼. 위시리스트와 넥스트 페스트로 알립니다.', '전 세계 액션, 인디 게이머', 70, INK),
             ('스토브 인디', '15%', '스마일게이트 발표(2023) 기준', '스마일게이트의 국내 인디 게임 플랫폼. 기획전과 커뮤니티가 있습니다.', '국내 인디 게임 이용자', 10, 'B9B3A8')]


def slide6():
    els = frame(6, '사업화', '사업화 계획 및 판매 채널')
    miles = [('2027년 2월', '데모 공개와 펀딩', '텀블벅 펀딩, 스팀 데모', '초기 수요와 암기 부담 검증'),
             ('2027년 8월', '얼리 액세스', '스팀, 스토브 인디 선판매', '판매 시작, 피드백 반영'),
             ('2028년 2월', '정식 출시', '5개 강토, 3개 결말', '10,000장 판매')]
    gap = 36; cw = (X1 - X0 - 2 * gap) / 3; y = 124
    for i, (d, t, sub, goal) in enumerate(miles):
        x = X0 + i * (cw + gap)
        if i:
            els.append(VR(x - gap / 2, y, 270))
        els.append(T(x, y, cw, 24, [[(d, BODY_B, SEAL if i == 0 else GREY)]], size=13))
        els.append(T(x, y + 26, cw, 30, [[(t, HEAD, INK)]], size=20))
        els.append(T(x, y + 64, cw, 22, sub, size=13.5))
        els.append(T(x, y + 90, cw, 22, [[('목표  ', BODY_B, INK), (goal, BODY, GREY)]], size=13))
    els.append(HR(X0, X1, 290, INK))
    pw = 250; cw2 = (X1 - X0 - pw - 3 * gap) / 3; y2 = 312
    for i, (name, fee, note, trait, cust, share, c) in enumerate(PLATFORMS):
        x = X0 + i * (cw2 + gap)
        if i:
            els.append(VR(x - gap / 2, y2, 660))
        els.append(T(x, y2, cw2, 26, [[(name, HEAD, INK)]], size=18))
        els.append(T(x, y2 + 36, cw2, 20, '예상 수수료', size=12, colour=GREY))
        els.append(T(x, y2 + 54, cw2, 44, [[(fee, HEAD, INK)]], size=34, lh=1.0))
        els.append(T(x, y2 + 92, cw2, 20, note, size=11, colour=GREY))
        els.append(T(x, y2 + 126, cw2, 70, trait, size=13, lh=1.55))
        els.append(T(x, y2 + 206, cw2, 44, [[('고객  ', BODY_B, INK), (cust, BODY, INK)]], size=13, lh=1.5))
    px = X1 - pw
    els.append(VR(px - gap / 2, y2, 660))
    els.append(T(px, y2, pw, 26, [[('예상 매출 비율', HEAD, INK)]], size=18))
    els.append(CHART('pie', px + 25, y2 + 40, 200, 200, cats=[p[0] for p in PLATFORMS], vals=[p[5] for p in PLATFORMS], colours=[p[6] for p in PLATFORMS]))
    yy = y2 + 256
    for p in sorted(PLATFORMS, key=lambda q: -q[5]):
        els.append(Rl(px, yy + 6, 10, 10, p[6]))
        els.append(T(px + 18, yy, 150, 22, p[0], size=13))
        els.append(T(px + 150, yy, pw - 150, 22, [[(f'{p[5]}%', BODY_B, INK)]], size=13, align='right'))
        yy += 26
    return els, ('2027년 2월 데모 공개와 텀블벅 펀딩으로 초기 수요를 확인하고, 8월 얼리 액세스를 거쳐 2028년 2월에 정식 출시합니다. '
                 '텀블벅에서는 후원자에게 먼저 팔고, 스팀에서는 전 세계 게이머에게, 스토브 인디에서는 국내 인디 게임 이용자에게 팝니다. 매출의 70%는 스팀에서 나올 것으로 봅니다.')


def slide7():
    els = frame(7, '팀역량', '팀 역량 및 보완 계획')
    gap = 48; cw = (X1 - X0 - gap) / 2; y = 124
    groups = [('대표 한예준', '기획, 개발, 아트', [('유학동양학 전공', '동양철학의 현대적 의미를 주제로 졸업논문을 썼습니다.'),
                                             ('게임 개발 부트캠프 수료', '수료 후 외주 개발로 실무를 익혔습니다.'),
                                             ('오행부 개발 현황', '5개 강토 월드, 한글 작도 인식, 전투와 요괴, 마석 자동차, 수묵 렌더링까지 플레이할 수 있는 빌드를 만들었습니다.')]),
              ('보완 계획', '2027년 1분기까지', [('사운드와 마케팅', 'AI로 시안을 만든 뒤 외주로 완성합니다.'), ('창업 프로그램', '멘토링과 펀딩으로 사업화 역량을 보강합니다.'),
                                           ('게임 행사', '행사와 전시에 나가 노출과 피드백을 얻습니다.')])]
    for i, (head, sub, items) in enumerate(groups):
        x = X0 + i * (cw + gap)
        if i:
            els.append(VR(x - gap / 2, y, 660))
        els.append(T(x, y, cw, 30, [[(head, HEAD, INK), ('   ' + sub, BODY, GREY)]], size=20))
        yy = y + 60
        for t, d in items:
            els.append(HR(x, x + cw, yy - 14))
            els.append(T(x, yy, cw, 26, [[(t, HEAD, INK)]], size=17))
            els.append(T(x, yy + 32, cw, 50, d, size=14, lh=1.55))
            yy += 118
    return els, ('저는 유학동양학을 전공하며 동양철학의 현대적 의미를 주제로 졸업논문을 썼고, 게임 개발 부트캠프를 수료한 뒤 외주 개발을 진행했습니다. '
                 '현재 오행부는 다섯 강토와 작도 인식, 전투, 수묵 렌더링까지 플레이할 수 있습니다. 부족한 사운드와 마케팅은 AI로 시안을 만든 뒤 외주로 완성하고, '
                 '창업 프로그램과 게임 행사를 활용해 2027년 1분기 안에 채우겠습니다.')


SLIDES = [slide1, slide2, slide3, slide4, slide5, slide6, slide7]


# ---------------------------------------------------------------- HTML board

def paras(runs, font, colour):
    if isinstance(runs, str):
        runs = [runs]
    out = []
    for p in runs:
        if isinstance(p, str):
            p = [(p, font, colour)]
        out.append(''.join(f'<span style="font-family:\'{f}\';color:#{c}">{html.escape(t)}</span>' for t, f, c in p))
    return out


def svg_chart(e):
    w, h = e['w'], e['h']
    if e['kind'] == 'line':
        vals, cats = e['vals'], e['cats']
        l, r, t, b = 44, 16, 24, 30
        sx = lambda i: l + i * (w - l - r) / (len(vals) - 1)
        sy = lambda v: t + (1 - v / 45) * (h - t - b)
        g = ''.join(f'<line x1="{l}" x2="{w - r}" y1="{sy(v):.1f}" y2="{sy(v):.1f}" stroke="#DCD8D0"/><text x="{l - 8}" y="{sy(v) + 4:.1f}" text-anchor="end">{v}%</text>' for v in range(0, 50, 10))
        path = ' '.join(f'{sx(i):.1f},{sy(v):.1f}' for i, v in enumerate(vals))
        pts = ''.join(f'<circle cx="{sx(i):.1f}" cy="{sy(v):.1f}" r="{6 if i == len(vals) - 1 else 3.5}" fill="#{SEAL if i == len(vals) - 1 else INK}"/>'
                      f'<text x="{sx(i):.1f}" y="{sy(v) - 12:.1f}" text-anchor="middle" style="font-weight:700;fill:#{SEAL if i == len(vals) - 1 else INK}">{v}%</text>' for i, v in enumerate(vals))
        cl = ''.join(f'<text x="{sx(i):.1f}" y="{h - 8}" text-anchor="middle">{c}</text>' for i, c in enumerate(cats))
        return f'<svg width="{w}" height="{h}" style="font:12px KoPubDotum Medium;fill:#{GREY}">{g}<polyline points="{path}" fill="none" stroke="#{INK}" stroke-width="2.5"/>{pts}{cl}</svg>'
    import math
    tot = sum(e['vals']); a0 = -math.pi / 2; cx = cy = w / 2; r = w / 2; out = ''
    for v, c in zip(e['vals'], e['colours']):
        a1 = a0 + 2 * math.pi * v / tot
        large = 1 if a1 - a0 > math.pi else 0
        out += f'<path d="M{cx},{cy} L{cx + r * math.cos(a0):.1f},{cy + r * math.sin(a0):.1f} A{r},{r} 0 {large} 1 {cx + r * math.cos(a1):.1f},{cy + r * math.sin(a1):.1f} Z" fill="#{c}" stroke="#{PAPER}" stroke-width="2"/>'
        a0 = a1
    return f'<svg width="{w}" height="{h}">{out}</svg>'


def render_html(pages):
    BOARD.mkdir(parents=True, exist_ok=True)
    faces = ''.join(f"@font-face{{font-family:'{n}';src:url('file:///C:/Windows/Fonts/{f}.ttf')}}" for n, f in
                    [(HEAD, 'KoPubBatangBold'), (HEAD_M, 'KoPubBatangMedium'), (BODY, 'KoPubDotumMedium'), (BODY_B, 'KoPubDotumBold')])
    out = [f'<!doctype html><html><head><meta charset="utf-8"><title>오행부 피칭 v5</title><style>{faces}'
           'body{margin:0;background:#8a8780;display:flex;flex-direction:column;align-items:center;gap:24px;padding:24px}'
           '.s{position:relative;width:1280px;height:720px;overflow:hidden}.s>*{position:absolute;box-sizing:border-box;margin:0}'
           'p{margin:0}</style></head><body>']
    for els, _ in pages:
        out.append('<div class="s">')
        for e in els:
            st = f"left:{e['x']:.1f}px;top:{e['y']:.1f}px;width:{e['w']:.1f}px;height:{e['h']:.1f}px" if 'x' in e else ''
            if e['k'] == 'rect':
                out.append(f'<div style="{st};background:#{e["fill"]}"></div>')
            elif e['k'] == 'text':
                ps = ''.join(f'<p>{p}</p>' for p in paras(e['runs'], e['font'], e['colour']))
                out.append(f'<div style="{st};font-size:{e["size"]}px;line-height:{e["lh"]};text-align:{e["align"]};font-family:\'{e["font"]}\';color:#{e["colour"]}">{ps}</div>')
            elif e['k'] == 'img':
                out.append(f'<img src="{e["path"].as_uri()}" style="{st};object-fit:cover">')
            elif e['k'] == 'video':
                out.append(f'<img src="{e["poster"].as_uri()}" style="{st};object-fit:cover">')
            elif e['k'] == 'poly':
                xs = [p[0] for p in e['points']]; ys = [p[1] for p in e['points']]
                x0, y0 = min(xs), min(ys)
                pts = ' '.join(f'{px - x0:.1f},{py - y0:.1f}' for px, py in e['points'])
                out.append(f'<svg style="left:{x0}px;top:{y0}px" width="{max(xs) - x0}" height="{max(ys) - y0}"><polygon points="{pts}" fill="#{e["fill"]}"/></svg>')
            elif e['k'] == 'chart':
                out.append(f'<div style="{st}">{svg_chart(e)}</div>')
        out.append('</div>')
    out.append('</body></html>')
    path = BOARD / 'board.html'
    path.write_text(''.join(out), encoding='utf-8')
    return path


# ---------------------------------------------------------------- PPTX

def px(v):
    return Emu(int(round(v * 9525)))            # 1 px at 96 dpi = 9525 EMU


def rgb(h):
    return RGBColor.from_string(h)


def run_font(r, font, size, colour):
    r.font.size = Pt(size * .75); r.font.color.rgb = rgb(colour); r.font.name = font
    rpr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {}); rpr.append(el)
        el.set('typeface', font)


def render_pptx(pages):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    for els, note in pages:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for e in els:
            if e['k'] == 'rect':
                sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e['fill']); sh.line.fill.background(); sh.shadow.inherit = False
            elif e['k'] == 'text':
                tb = s.shapes.add_textbox(px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                runs = e['runs'] if not isinstance(e['runs'], str) else [e['runs']]
                for i, p in enumerate(runs):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = {'left': PP_ALIGN.LEFT, 'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}[e['align']]
                    para.line_spacing = e['lh'] * .83          # CSS line-height -> PowerPoint multiple (font ascent/descent differ)
                    for t, f, c in ([(p, e['font'], e['colour'])] if isinstance(p, str) else p):
                        r = para.add_run(); r.text = t; run_font(r, f, e['size'], c)
            elif e['k'] == 'img':
                iw, ih = Image.open(e['path']).size
                pic = s.shapes.add_picture(str(e['path']), px(e['x']), px(e['y']), px(e['w']), px(e['h']))
                if iw / ih > e['w'] / e['h']:
                    c = (1 - (e['w'] / e['h']) / (iw / ih)) / 2; pic.crop_left = pic.crop_right = c
                else:
                    c = (1 - (iw / ih) / (e['w'] / e['h'])) / 2; pic.crop_top = pic.crop_bottom = c
            elif e['k'] == 'video':
                s.shapes.add_movie(str(e['path']), px(e['x']), px(e['y']), px(e['w']), px(e['h']), poster_frame_image=str(e['poster']), mime_type='video/mp4')
            elif e['k'] == 'poly':
                pts = e['points']
                fb = s.shapes.build_freeform(px(pts[0][0]), px(pts[0][1]), scale=1.0)
                fb.add_line_segments([(px(a), px(b)) for a, b in pts[1:]], close=True)
                sh = fb.convert_to_shape()
                sh.fill.solid(); sh.fill.fore_color.rgb = rgb(e['fill']); sh.line.fill.background(); sh.shadow.inherit = False
            elif e['k'] == 'chart':
                data = CategoryChartData(); data.categories = e['cats']
                if e['kind'] == 'line':
                    data.add_series('마진율', e['vals'])
                    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, px(e['x']), px(e['y']), px(e['w']), px(e['h']), data).chart
                    ch.has_title = False; ch.has_legend = False
                    ch.font.size = Pt(9); ch.font.color.rgb = rgb(GREY); ch.font.name = BODY
                    plot = ch.plots[0]; plot.has_data_labels = True
                    dl = plot.data_labels; dl.number_format = '0.0"%"'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.ABOVE
                    dl.font.size = Pt(10); dl.font.bold = True; dl.font.color.rgb = rgb(INK)
                    ser = plot.series[0]; ser.format.line.color.rgb = rgb(INK); ser.format.line.width = Pt(2); ser.smooth = False
                    ser.marker.style = 8; ser.marker.size = 6
                    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(INK); ser.marker.format.line.color.rgb = rgb(INK)
                    sdl = ser.data_labels                   # series-level labels (a point override otherwise hides the rest)
                    sdl.show_value = True; sdl.number_format = '0.0"%"'; sdl.number_format_is_linked = False
                    sdl.position = XL_LABEL_POSITION.ABOVE; sdl.font.size = Pt(10); sdl.font.bold = True; sdl.font.color.rgb = rgb(INK)
                    last = ser.points[len(e['vals']) - 1]
                    last.marker.size = 10; last.marker.format.fill.solid(); last.marker.format.fill.fore_color.rgb = rgb(SEAL); last.marker.format.line.color.rgb = rgb(SEAL)
                    last.data_label.font.color.rgb = rgb(SEAL); last.data_label.font.size = Pt(12); last.data_label.font.bold = True
                    last.data_label.position = XL_LABEL_POSITION.ABOVE
                    lbl = last.data_label._dLbl
                    lbl.insert(1, lbl.makeelement(qn('c:numFmt'), {'formatCode': '0.0"%"', 'sourceLinked': '0'}))
                    va = ch.value_axis; va.maximum_scale = 45; va.minimum_scale = 0; va.major_unit = 10
                    va.major_gridlines.format.line.color.rgb = rgb('DCD8D0'); va.major_gridlines.format.line.width = Pt(.5)
                    va.format.line.fill.background(); va.tick_labels.number_format = '0"%"'; va.tick_labels.number_format_is_linked = False
                    ca = ch.category_axis; ca.format.line.color.rgb = rgb(RULE)
                else:
                    data.add_series('비율', [v / 100 for v in e['vals']])
                    ch = s.shapes.add_chart(XL_CHART_TYPE.PIE, px(e['x']), px(e['y']), px(e['w']), px(e['h']), data).chart
                    ch.has_title = False; ch.has_legend = False
                    plot = ch.plots[0]; plot.has_data_labels = False
                    for i, c in enumerate(e['colours']):
                        pt = plot.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
                        pt.format.line.color.rgb = rgb(PAPER); pt.format.line.width = Pt(1.5)
        s.notes_slide.notes_text_frame.text = note
    prs.save(DECK)
    if '--no-embed' not in sys.argv:
        ps = (f'$app = New-Object -ComObject PowerPoint.Application; $p = $app.Presentations.Open("{DECK}", $false, $false, $false); '
              f'$p.SaveAs("{DECK}", 24, -1); $p.Close()')
        subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True)
    return DECK


def main():
    BOARD.mkdir(parents=True, exist_ok=True)
    pages = [f() for f in SLIDES]
    print(render_html(pages))
    if '--html-only' not in sys.argv:
        print(render_pptx(pages))


if __name__ == '__main__':
    main()
